"""
PriceHunter — автоматический подбор лучших товаров на Wildberries.
Запускается в GitHub Actions (repository_dispatch), а не на серверах Make,
т.к. Wildberries блокирует запросы с IP облачных платформ типа Make.com.

Логика подбора товаров взята из пользовательского скрипта
wildberries_parser_on_catalog.py и адаптирована для запуска без интерактивного
ввода — все параметры приходят через переменные окружения (их передаёт
Make через repository_dispatch, GitHub Actions прокидывает их в env).

Как ищем товар (два уровня):
  1) Основной способ — настоящий полнотекстовый поиск Wildberries
     (search.wb.ru), тот же самый движок, что работает в строке поиска
     на сайте. Не привязан к статичному дереву категорий, поэтому находит
     любой товар точно, как и сам сайт.
  2) Запасной способ — если Wildberries не пропускает полнотекстовый поиск
     для внешних запросов, подбираем ближайшую по названию категорию
     каталога (catalog.wb.ru) и ищем в ней. Менее точно, но работает
     стабильно в любом случае.

Переменные окружения:
    QUERY               - поисковый запрос пользователя, например "велосипед спортивный"
    MIN_PRICE           - минимальная цена (руб)
    MAX_PRICE           - максимальная цена (руб)
    CHAT_ID             - id чата в Telegram, куда прислать результат
    TELEGRAM_BOT_TOKEN  - токен бота (секрет репозитория, НЕ хардкодить)
    PAGES               - сколько страниц каталога смотреть (по умолчанию 3)
    TOP_N               - сколько лучших товаров показать в сообщении (по умолчанию 5)
"""

import os
import re
import time
import random
import json
import difflib
import urllib.parse
import requests
import pandas as pd
import pymorphy3
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

WB_DEST_MOSCOW = -1257786  # код региона "Москва и область" в API Wildberries

_morph = pymorphy3.MorphAnalyzer()


def _lemmatize(word: str) -> str:
    """Приводим слово к начальной форме (лемме), чтобы "чехлы" и "чехол",
    "велосипеды" и "велосипед" и т.п. считались одним и тем же словом."""
    word = word.replace("ё", "е")
    try:
        return _morph.parse(word)[0].normal_form.replace("ё", "е")
    except Exception:
        return word


def _wb_headers() -> dict:
    return {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:109.0) AppleWebKit/537.36",
        "Accept": "application/json",
        "Accept-Language": "ru-RU,ru;q=0.9,en-US;q=0.8,en;q=0.7",
        "Connection": "keep-alive",
        "Referer": "https://www.wildberries.ru/",
    }


def search_wb_fulltext(query: str, page: int, low_price: int, top_price: int, retry_count: int = 0):
    """
    Настоящий полнотекстовый поиск Wildberries (тот же движок, что и строка
    поиска на сайте) — ищет напрямую по тексту запроса, без привязки к
    статичному дереву категорий. Возвращает None, если Wildberries не
    пропустил запрос (сайт может заблокировать полнотекстовый поиск для
    запросов не со своего сайта) — тогда main() переключается на поиск
    по категориям.
    """
    url = (
        "https://search.wb.ru/exactmatch/ru/common/v13/search"
        f"?appType=1&curr=rub&dest={WB_DEST_MOSCOW}&locale=ru"
        f"&page={page}&priceU={low_price * 100};{top_price * 100}"
        f"&sort=rate&spp=0&suppressSpellcheck=false&resultset=catalog"
        f"&query={urllib.parse.quote(query)}"
    )
    time.sleep(random.uniform(2.0, 4.0) + retry_count * 2)
    try:
        r = requests.get(url, headers=_wb_headers(), timeout=15)
        print(f"[DEBUG search] status={r.status_code} url={url} body={r.text[:300]}")
    except requests.exceptions.RequestException as e:
        if retry_count < 3:
            time.sleep(5)
            return search_wb_fulltext(query, page, low_price, top_price, retry_count + 1)
        return None
    if r.status_code == 429 and retry_count < 3:
        time.sleep(20 + retry_count * 10)
        return search_wb_fulltext(query, page, low_price, top_price, retry_count + 1)
    if r.status_code != 200:
        return None
    try:
        return r.json()
    except json.JSONDecodeError:
        return None


def get_catalogs_wb() -> dict:
    """Полный каталог категорий Wildberries (статический CDN-файл, не банится)."""
    url = "https://static-basket-01.wbbasket.ru/vol0/data/main-menu-ru-ru-v3.json"
    headers = {"Accept": "*/*", "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}
    r = requests.get(url, headers=headers, timeout=20)
    r.raise_for_status()
    return r.json()


def get_data_category(catalogs_wb) -> list:
    """Разворачиваем дерево категорий в плоский список {name, shard, url, query}."""
    catalog_data = []
    stack = [catalogs_wb]
    while stack:
        current = stack.pop()
        if isinstance(current, dict):
            if "childs" not in current:
                if current.get("shard") and current.get("query"):
                    catalog_data.append({
                        "name": current["name"],
                        "shard": current.get("shard"),
                        "url": current.get("url"),
                        "query": current.get("query"),
                    })
            else:
                stack.append(current["childs"])
        elif isinstance(current, list):
            for item in reversed(current):
                stack.append(item)
    return catalog_data


def find_category_by_query(query: str, catalog_list: list):
    """
    Запасной способ (используется, только если search_wb_fulltext
    недоступен): ищем наиболее подходящую категорию каталога по свободному
    тексту запроса. Сначала точное совпадение по леммам (начальным формам
    слов), затем — если точного нет — ближайшая по написанию категория, но
    ТОЛЬКО если сходство действительно высокое. Если ничего убедительного
    не нашлось — возвращаем None и честно говорим пользователю, что не
    нашли, а не подсовываем случайную категорию не по теме.
    """
    STOPWORDS = {"для", "или", "как", "что", "это", "она", "мой", "моя", "при", "под", "над", "без", "про", "чем", "уже", "если", "все", "всё", "чтобы", "когда", "куда", "этот", "эта", "эти"}
    raw_words = [w.lower() for w in re.split(r"[\s,]+", query.strip()) if len(w) > 2 and w.lower() not in STOPWORDS]
    if not raw_words:
        return None

    query_lemmas = {_lemmatize(w) for w in raw_words}

    # 1) точное совпадение по леммам
    best_match, best_score = None, 0
    for catalog in catalog_list:
        name = catalog.get("name", "").lower()
        name_words = re.split(r"[^а-яё]+", name)
        name_lemmas = {_lemmatize(nw) for nw in name_words if len(nw) > 2}
        score = len(query_lemmas & name_lemmas)
        if score > best_score:
            best_score = score
            best_match = catalog
    if best_score > 0:
        return best_match

    # 2) запасной вариант: ближайшая по написанию категория, но только при
    # высоком сходстве (порог подобран так, чтобы не путать случайные слова
    # вроде "отвертка"/"подсветка" — лучше честно не найти, чем ошибиться)
    best_match, best_ratio = None, 0.0
    for catalog in catalog_list:
        name = catalog.get("name", "").lower()
        name_words = [w for w in re.split(r"[^а-яё]+", name) if len(w) > 2]
        for qw in query_lemmas:
            for nw in name_words:
                ratio = difflib.SequenceMatcher(None, qw, nw).ratio()
                if ratio > best_ratio:
                    best_ratio = ratio
                    best_match = catalog

    if best_ratio >= 0.82:
        return best_match
    return None


def get_data_from_json(json_file: dict) -> list:
    data_list = []
    json_file = json_file or {}
    products = json_file.get("products")
    if products is None:
        products = json_file.get("data", {}).get("products", [])
    for data in products:
        try:
            sku = data.get("id")
            name = data.get("name", "")
            try:
                if data.get("sizes"):
                    price = int(data["sizes"][0].get("price", {}).get("product", 0) / 100)
                    basic = int(data["sizes"][0].get("price", {}).get("basic", 0) / 100)
                else:
                    price = int(data.get("salePriceU", 0) / 100) if data.get("salePriceU") else 0
                    basic = int(data.get("priceU", 0) / 100) if data.get("priceU") else price
                discount_percent = round((1 - price / basic) * 100, 1) if basic > 0 and price > 0 else 0
            except (KeyError, TypeError, IndexError, ZeroDivisionError):
                price, basic, discount_percent = 0, 0, 0

            data_list.append({
                "ID": sku,
                "Название": name,
                "Бренд": data.get("brand", ""),
                "Цена без скидки": basic,
                "Цена со скидкой": price,
                "Скидка %": discount_percent,
                "Рейтинг товара": data.get("rating", 0),
                "Рейтинг продавца": data.get("supplierRating", 0),
                "Продавец": data.get("supplier", ""),
                "Количество отзывов": data.get("feedbacks", 0),
                "Рейтинг отзывов": data.get("reviewRating", 0),
                "Ссылка": f"https://www.wildberries.ru/catalog/{sku}/detail.aspx?targetUrl=BP",
            })
        except Exception:
            continue
    return data_list


def scrap_page(page: int, shard: str, query: str, low_price: int, top_price: int, retry_count: int = 0) -> dict:
    headers = _wb_headers()
    url = (
        f"https://catalog.wb.ru/catalog/{shard}/v4/catalog"
        f"?appType=1&curr=rub&dest={WB_DEST_MOSCOW}&locale=ru"
        f"&page={page}&priceU={low_price * 100};{top_price * 100}"
        f"&sort=rate&spp=0&{query}"
    )
    time.sleep(random.uniform(2.0, 4.0) + retry_count * 2)
    try:
        r = requests.get(url, headers=headers, timeout=15)
        print(f"[DEBUG] status={r.status_code} url={url} body={r.text[:500]}")
    except requests.exceptions.RequestException as e:
        if retry_count < 3:
            time.sleep(5)
            return scrap_page(page, shard, query, low_price, top_price, retry_count + 1)
        raise
    if r.status_code == 429 and retry_count < 3:
        time.sleep(20 + retry_count * 10)
        return scrap_page(page, shard, query, low_price, top_price, retry_count + 1)
    if r.status_code != 200:
        return {"data": {"products": []}}
    try:
        return r.json()
    except json.JSONDecodeError:
        return {"data": {"products": []}}


def save_excel(data: list, filepath: str):
    df = pd.DataFrame(data)
    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="data", index=False)
        worksheet = writer.sheets["data"]

        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(left=Side(style="thin"), right=Side(style="thin"),
                              top=Side(style="thin"), bottom=Side(style="thin"))

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border

        for column in worksheet.columns:
            max_length = max((len(str(c.value)) for c in column if c.value is not None), default=0)
            worksheet.column_dimensions[get_column_letter(column[0].column)].width = min(max_length + 2, 60)

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions


def send_telegram_message(token: str, chat_id: str, text: str):
    requests.post(
        f"https://api.telegram.org/bot{token}/sendMessage",
        json={"chat_id": chat_id, "text": text, "disable_web_page_preview": True},
        timeout=20,
    )


def send_telegram_document(token: str, chat_id: str, filepath: str, caption: str = ""):
    with open(filepath, "rb") as f:
        requests.post(
            f"https://api.telegram.org/bot{token}/sendDocument",
            data={"chat_id": chat_id, "caption": caption},
            files={"document": f},
            timeout=60,
        )


def main():
    query = os.environ["QUERY"]
    min_price = int(os.environ.get("MIN_PRICE") or 1)
    max_price = int(os.environ.get("MAX_PRICE") or 10_000_000)
    chat_id = os.environ["CHAT_ID"]
    token = os.environ["TELEGRAM_BOT_TOKEN"]
    pages = int(os.environ.get("PAGES") or 3)
    top_n = int(os.environ.get("TOP_N") or 5)

    try:
        data_list = []
        category_name = None

        # 1) основной способ — настоящий полнотекстовый поиск Wildberries
        search_available = True
        empty_pages = 0
        for page in range(1, pages + 1):
            result = search_wb_fulltext(query, page, min_price, max_price)
            if result is None:
                search_available = False
                break
            products = get_data_from_json(result)
            if not products:
                empty_pages += 1
                if empty_pages >= 2:
                    break
                continue
            data_list.extend(products)

        # 2) запасной способ — если Wildberries не пропустил полнотекстовый
        # поиск, ищем по ближайшей категории каталога
        if not search_available:
            data_list = []
            catalog_list = get_data_category(get_catalogs_wb())
            category = find_category_by_query(query, catalog_list)

            if category is None:
                send_telegram_message(
                    token, chat_id,
                    f"Не смог найти подходящую категорию на Wildberries для запроса «{query}». "
                    f"Попробуй сформулировать проще (например, одно-два слова)."
                )
                return

            category_name = category["name"]
            empty_pages = 0
            for page in range(1, pages + 1):
                result = scrap_page(page, category["shard"], category["query"], min_price, max_price)
                products = get_data_from_json(result)
                if not products:
                    empty_pages += 1
                    if empty_pages >= 2:
                        break
                    continue
                data_list.extend(products)

        if not data_list:
            where = f" в категории «{category_name}»" if category_name else ""
            send_telegram_message(
                token, chat_id,
                f"По запросу «{query}»{where} ничего не нашлось "
                f"в диапазоне {min_price}–{max_price} ₽. Попробуй расширить диапазон цен."
            )
            return

        # автоподбор лучших: сортировка по рейтингу товара, затем по количеству отзывов
        data_list.sort(key=lambda p: (p["Рейтинг товара"], p["Количество отзывов"]), reverse=True)

        top_products = data_list[:top_n]
        title_suffix = f" ({category_name})" if category_name else ""
        lines = [f"Готово! 🎉 Нашёл и отсортировал товары по «{query}»{title_suffix} от {min_price} до {max_price} ₽.\n"]
        lines.append(f"🏆 Топ-{len(top_products)} по рейтингу и отзывам:\n")
        for i, p in enumerate(top_products, 1):
            lines.append(
                f"{i}. {p['Название']}\n"
                f"   💰 {p['Цена со скидкой']} ₽ | ⭐ {p['Рейтинг товара']} "
                f"({p['Количество отзывов']} отзывов)\n"
                f"   {p['Ссылка']}\n"
            )
        lines.append(f"\nПолный список из {len(data_list)} товаров — во вложенном Excel-файле.")
        send_telegram_message(token, chat_id, "\n".join(lines))

        os.makedirs("out", exist_ok=True)
        safe_name = re.sub(r"[^\w\-]+", "_", query)[:40]
        filepath = f"out/{safe_name}.xlsx"
        save_excel(data_list, filepath)
        send_telegram_document(token, chat_id, filepath, caption=f"Все товары по запросу «{query}»")

    except Exception as e:
        send_telegram_message(token, chat_id, f"Произошла ошибка при поиске: {e}")
        raise


if __name__ == "__main__":
    main()
